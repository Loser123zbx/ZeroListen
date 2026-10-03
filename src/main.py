# -*- coding: utf-8 -*-
"""零听 ZeroListen —— 本地离线中/英文词句 TTS 批量生成工具。

入口: python main.py  (在 src 目录下运行)

说明:
- GUI 面板由 wxFormBuilder 生成于 all_panels.py (请勿编辑)。
- 本文件负责把面板接线、实现词句库管理、TTS 生成、批量导出、
  进度展示, 以及"音频导出配置"对话框(该界面为后补实现)。
- 语音合成走本地 kokoro-js (Node), 首次运行会从 hf-mirror 下载并缓存模型。
"""

import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import zipfile
import urllib.request
import urllib.error

import wx
import wx.adv
import wx.grid
from openpyxl import Workbook, load_workbook

from all_panels import main_panel, wordlib_grid, welcome_page, workbar_page, progress
from html_player import render_player_html

def _app_dir():
    # 打包成可执行文件时, 资源(tts_cli.js/node_modules/模型)放在 exe 同级目录。
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


APP_DIR = _app_dir()
PROJECT_ROOT = os.path.dirname(APP_DIR)
NODE_DIR = os.path.join(PROJECT_ROOT, "node")
LOG_DIR = os.path.join(PROJECT_ROOT, "log")
WORDLIB_DIR = os.path.join(APP_DIR, "wordlibs")
EXPORT_DIR = os.path.join(APP_DIR, "exports")
CONFIG_PATH = os.path.join(APP_DIR, "config.json")
TTS_CLI = os.path.join(APP_DIR, "tts_cli.js")

os.makedirs(LOG_DIR, exist_ok=True)


def _apply_button_style(button):
    button.SetFont(wx.Font(9, wx.FONTFAMILY_DEFAULT, wx.FONTSTYLE_NORMAL, wx.FONTWEIGHT_BOLD, False))
    button.SetForegroundColour(wx.Colour(255, 255, 255))
    button.SetBackgroundColour(wx.Colour(25, 25, 112))
    return button


DEFAULT_MODEL_ID = "onnx-community/Kokoro-82M-ONNX"
DEFAULT_TEST_TEXT = "The universe said I love you because you are love."

# 语音列表: (voice_id, 显示名)。kokoro-js 内置全部 voice。
VOICES = [
    ("af_heart", "af_heart（美音·女声）"),
    ("af_alloy", "af_alloy（美音·女声）"),
    ("af_aoede", "af_aoede（美音·女声）"),
    ("af_bella", "af_bella（美音·女声）"),
    ("af_jessica", "af_jessica（美音·女声）"),
    ("af_kore", "af_kore（美音·女声）"),
    ("af_nicole", "af_nicole（美音·女声）"),
    ("af_nova", "af_nova（美音·女声）"),
    ("af_river", "af_river（美音·女声）"),
    ("af_sarah", "af_sarah（美音·女声）"),
    ("af_sky", "af_sky（美音·女声）"),
    ("am_adam", "am_adam（美音·男声）"),
    ("am_echo", "am_echo（美音·男声）"),
    ("am_eric", "am_eric（美音·男声）"),
    ("am_fenrir", "am_fenrir（美音·男声）"),
    ("am_liam", "am_liam（美音·男声）"),
    ("am_michael", "am_michael（美音·男声）"),
    ("am_onyx", "am_onyx（美音·男声）"),
    ("am_puck", "am_puck（美音·男声）"),
    ("am_santa", "am_santa（美音·男声）"),
    ("bf_emma", "bf_emma（英音·女声）"),
    ("bf_isabella", "bf_isabella（英音·女声）"),
    ("bf_alice", "bf_alice（英音·女声）"),
    ("bf_lily", "bf_lily（英音·女声）"),
    ("bm_george", "bm_george（英音·男声）"),
    ("bm_lewis", "bm_lewis（英音·男声）"),
    ("bm_daniel", "bm_daniel（英音·男声）"),
    ("bm_fable", "bm_fable（英音·男声）"),
    ("zf_xiaobei", "zf_xiaobei（普通话·女声）"),
    ("zf_xiaoni", "zf_xiaoni（普通话·女声）"),
    ("zf_xiaoxiao", "zf_xiaoxiao（普通话·女声）"),
    ("zf_xiaoyi", "zf_xiaoyi（普通话·女声）"),
    ("zm_yunjian", "zm_yunjian（普通话·男声）"),
    ("zm_yunxi", "zm_yunxi（普通话·男声）"),
    ("zm_yunxia", "zm_yunxia（普通话·男声）"),
    ("zm_yunyang", "zm_yunyang（普通话·男声）"),
]

class AudioConfig(object):
    """音频导出配置(含 HTML 播放器跟读默认设置、窗口视觉配置)。"""

    def __init__(self, d=None):
        d = d or {}
        self.voice = d.get("voice", "af_heart")
        self.speed = float(d.get("speed", 1.0))
        self.dtype = d.get("dtype", "fp32")
        self.repeat_count = int(d.get("repeat_count", 2))
        self.interval_mode = d.get("interval_mode", "seconds")  # seconds | multiple
        self.interval_value = float(d.get("interval_value", 1.5))
        self.bg_color = self._normalize_hex(d.get("bg_color", "#FFFFFF"), "#FFFFFF")
        self.accent_color = self._normalize_hex(d.get("accent_color", "#191970"), "#191970")
        self.window_alpha = int(d.get("window_alpha", 255))
        self._normalize()

    @staticmethod
    def _normalize_hex(value, default):
        if not isinstance(value, str):
            return default
        value = value.strip()
        if value.startswith("#"):
            value = value[1:]
        if len(value) != 6 or any(ch not in "0123456789abcdefABCDEF" for ch in value):
            return default
        return "#" + value.upper()

    def _normalize(self):
        if self.speed < 0.5:
            self.speed = 0.5
        if self.speed > 2.0:
            self.speed = 2.0
        if self.dtype not in ("fp32", "fp16", "q8", "q4", "q4f16"):
            self.dtype = "fp32"
        if self.repeat_count < 1:
            self.repeat_count = 1
        if self.repeat_count > 10:
            self.repeat_count = 10
        if self.interval_mode not in ("seconds", "multiple"):
            self.interval_mode = "seconds"
        if self.interval_value < 0:
            self.interval_value = 0
        self.bg_color = self._normalize_hex(self.bg_color, "#FFFFFF")
        self.accent_color = self._normalize_hex(self.accent_color, "#191970")
        self.window_alpha = max(20, min(255, int(self.window_alpha)))

    def to_dict(self):
        return {
            "voice": self.voice,
            "speed": self.speed,
            "dtype": self.dtype,
            "repeat_count": self.repeat_count,
            "interval_mode": self.interval_mode,
            "interval_value": self.interval_value,
            "bg_color": self.bg_color,
            "accent_color": self.accent_color,
            "window_alpha": self.window_alpha,
        }


def load_config():
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            return AudioConfig(json.load(f))
    except Exception:
        return AudioConfig()


def save_config(cfg):
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg.to_dict(), f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def hex_to_rgb(color_hex):
    color_hex = str(color_hex or "#FFFFFF").strip()
    if color_hex.startswith("#"):
        color_hex = color_hex[1:]
    if len(color_hex) == 6:
        try:
            return tuple(int(color_hex[i:i+2], 16) for i in (0, 2, 4))
        except ValueError:
            pass
    return (255, 255, 255)


def apply_visual_theme(frame, config):
    bg_color = wx.Colour(*hex_to_rgb(config.bg_color))
    accent = wx.Colour(*hex_to_rgb(config.accent_color))
    frame.SetBackgroundColour(bg_color)
    frame.SetTransparent(int(config.window_alpha))
    for child in getattr(frame, "GetChildren", lambda: [])():
        try:
            child.SetBackgroundColour(bg_color)
        except Exception:
            pass
    if hasattr(frame, "workbar"):
        frame.workbar.SetBackgroundColour(bg_color)
    if hasattr(frame, "welcome"):
        frame.welcome.SetBackgroundColour(bg_color)
    if hasattr(frame, "editor"):
        frame.editor.SetBackgroundColour(bg_color)
    if hasattr(frame, "exporter"):
        frame.exporter.SetBackgroundColour(bg_color)
    return accent


def write_error_log(message, exc=None):
    os.makedirs(LOG_DIR, exist_ok=True)
    now = datetime.datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    log_path = os.path.join(LOG_DIR, f"error_{now}.log")
    details = str(message or "")
    if exc is not None:
        details += "\n\n" + traceback.format_exc() if exc is True else "\n\n" + str(exc)
    with open(log_path, "w", encoding="utf-8") as f:
        f.write(details)
    return log_path


def show_error_dialog(title, message, exc=None):
    details = str(message or "")
    if exc is not None:
        if exc is True:
            details += "\n\n" + traceback.format_exc()
        else:
            details += "\n\n" + str(exc)
    log_path = write_error_log(details)
    details += "\n\n日志已保存：\n" + log_path

    dlg = wx.Dialog(None, title=title, size=(700, 420))

    main_sizer = wx.BoxSizer(wx.VERTICAL)
    text = wx.TextCtrl(dlg, value=details, style=wx.TE_MULTILINE | wx.TE_READONLY | wx.TE_RICH2)
    text.SetInsertionPoint(0)
    main_sizer.Add(text, 1, wx.EXPAND | wx.ALL, 10)

    btn_sizer = wx.BoxSizer(wx.HORIZONTAL)
    copy_btn = wx.Button(dlg, label="复制错误")
    open_btn = wx.Button(dlg, label="打开日志目录")
    close_btn = wx.Button(dlg, label="关闭")

    def on_copy(evt):
        try:
            if wx.TheClipboard.Open():
                wx.TheClipboard.SetData(wx.TextDataObject(details))
                wx.TheClipboard.Close()
                wx.MessageBox("错误信息已复制到剪贴板。", "已复制", wx.OK | wx.ICON_INFORMATION)
            else:
                wx.MessageBox("剪贴板不可用，请手动复制下方内容。", "提示", wx.OK | wx.ICON_WARNING)
        except Exception:
            wx.MessageBox("复制失败，请手动复制下方内容。", "提示", wx.OK | wx.ICON_WARNING)

    def on_open_log(evt):
        try:
            os.startfile(LOG_DIR)
        except Exception:
            wx.MessageBox("日志目录：\n" + LOG_DIR, "日志位置", wx.OK | wx.ICON_INFORMATION)

    copy_btn.Bind(wx.EVT_BUTTON, on_copy)
    open_btn.Bind(wx.EVT_BUTTON, on_open_log)
    close_btn.Bind(wx.EVT_BUTTON, lambda evt: dlg.EndModal(wx.ID_OK))

    btn_sizer.Add(copy_btn, 0, wx.RIGHT | wx.ALIGN_CENTER_VERTICAL, 8)
    btn_sizer.Add(open_btn, 0, wx.RIGHT | wx.ALIGN_CENTER_VERTICAL, 8)
    btn_sizer.Add(close_btn, 0, wx.ALIGN_CENTER_VERTICAL)
    main_sizer.Add(btn_sizer, 0, wx.ALIGN_RIGHT | wx.ALL, 10)

    dlg.SetSizerAndFit(main_sizer)
    dlg.ShowModal()
    dlg.Destroy()
    return log_path


def install_global_exception_handler():
    def _handler(exc_type, exc_value, exc_tb):
        message = "".join(traceback.format_exception(exc_type, exc_value, exc_tb))
        show_error_dialog("程序异常", str(exc_value) or exc_type.__name__, message)
    sys.excepthook = _handler


# ---------------------------------------------------------------------------
# 词句库文件层
# ---------------------------------------------------------------------------

def _find_node_executable(start_dir):
    if not start_dir or not os.path.isdir(start_dir):
        return None
    for root, _dirs, files in os.walk(start_dir):
        for name in files:
            low = name.lower()
            if low in ("node.exe", "node"):
                return os.path.join(root, name)
    return None


def _download_node_to_local_dir(prog_cb=None):
    os.makedirs(NODE_DIR, exist_ok=True)
    if os.name != "nt":
        return None
    existing = _find_node_executable(NODE_DIR)
    if existing:
        return existing
    try:
        if prog_cb is not None:
            prog_cb(0, 100, "正在获取 Node.js 下载信息…")
        with urllib.request.urlopen("https://nodejs.org/dist/index.json", timeout=30) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        release = next((item for item in data if item.get("lts") not in (None, False)), data[0] if data else None)
        if not release:
            return None
        version = release.get("version")
        if not version:
            return None
        zip_url = f"https://nodejs.org/dist/{version}/node-{version}-win-x64.zip"
        zip_path = os.path.join(NODE_DIR, "node-win-x64.zip")
        if prog_cb is not None:
            prog_cb(5, 100, f"正在下载 Node.js {version}…")
        with urllib.request.urlopen(zip_url, timeout=60) as resp, open(zip_path, "wb") as f:
            total = int(resp.headers.get("Content-Length", "0") or "0")
            received = 0
            while True:
                chunk = resp.read(65536)
                if not chunk:
                    break
                f.write(chunk)
                received += len(chunk)
                if total > 0 and prog_cb is not None:
                    prog_cb(min(95, int(received * 100 / total)), 100, f"正在下载 Node.js {version}… {int(received * 100 / total)}%")
        if prog_cb is not None:
            prog_cb(96, 100, "正在解压 Node.js…")
        with zipfile.ZipFile(zip_path, "r") as zf:
            zf.extractall(NODE_DIR)
        os.remove(zip_path)
        if prog_cb is not None:
            prog_cb(100, 100, "Node.js 下载完成")
        return _find_node_executable(NODE_DIR)
    except Exception:
        return None


def _resolve_node_executable(prog_cb=None):
    candidates = [
        os.path.join(NODE_DIR, "node.exe"),
        os.path.join(NODE_DIR, "node"),
        os.path.join(APP_DIR, "node.exe"),
        os.path.join(APP_DIR, "node"),
        os.path.join(PROJECT_ROOT, "node.exe"),
        os.path.join(PROJECT_ROOT, "node"),
    ]
    for item in candidates:
        if os.path.isfile(item):
            return item
    found = _find_node_executable(NODE_DIR)
    if found:
        return found
    found = shutil.which("node") or shutil.which("node.exe")
    if found:
        return found
    found = _download_node_to_local_dir(prog_cb)
    return found


def _resolve_npm_command(node_exe):
    if not node_exe:
        return None
    base_dir = os.path.dirname(node_exe)
    for candidate in [
        os.path.join(base_dir, "npm.cmd"),
        os.path.join(base_dir, "npm"),
        os.path.join(base_dir, "npx.cmd"),
        os.path.join(base_dir, "npx"),
    ]:
        if os.path.isfile(candidate):
            return candidate
    return None


def _node_env_with_local_path(node_exe):
    env = os.environ.copy()
    if not node_exe:
        return env
    node_dir = os.path.dirname(node_exe)
    path_value = env.get("PATH", "")
    env["PATH"] = node_dir + os.pathsep + path_value if path_value else node_dir
    return env


def _ensure_node_dependencies(node_exe, prog_cb=None):
    package_json = os.path.join(APP_DIR, "package.json")
    if not os.path.exists(package_json):
        return
    node_modules = os.path.join(APP_DIR, "node_modules")
    kokoro_js = os.path.join(node_modules, "kokoro-js")
    transformers = os.path.join(node_modules, "@huggingface", "transformers")
    package_ok = os.path.isdir(kokoro_js) and os.path.exists(os.path.join(kokoro_js, "package.json"))
    package_ok = package_ok and os.path.isdir(transformers) and os.path.exists(os.path.join(transformers, "package.json"))
    if package_ok:
        return

    npm_cmd = _resolve_npm_command(node_exe)
    if not npm_cmd:
        return

    if os.path.isdir(node_modules):
        shutil.rmtree(node_modules, ignore_errors=True)

    for attempt in range(2):
        if prog_cb is not None:
            prog_cb(0, 100, "正在检查并安装项目依赖…")
        cmd = [npm_cmd, "install", "--prefix", APP_DIR, "--no-fund", "--no-audit"]
        proc = subprocess.Popen(cmd, cwd=APP_DIR, stdout=subprocess.PIPE,
                                stderr=subprocess.STDOUT, text=True,
                                encoding="utf-8", errors="replace",
                                env=_node_env_with_local_path(node_exe))
        tail = []
        try:
            for line in proc.stdout:
                line = line.rstrip("\n")
                if line:
                    tail.append(line)
                    if len(tail) > 30:
                        tail.pop(0)
                    if prog_cb is not None:
                        prog_cb(10, 100, line.strip()[:80] or "正在安装项目依赖…")
            rc = proc.wait()
        finally:
            if proc.stdout:
                proc.stdout.close()

        if prog_cb is not None:
            prog_cb(100, 100, "依赖安装完成" if rc == 0 else "依赖安装失败")
        if rc == 0:
            break
        if os.path.isdir(node_modules):
            shutil.rmtree(node_modules, ignore_errors=True)
        if attempt == 0:
            continue
        raise RuntimeError("自动安装项目依赖失败：\n" + ("\n".join(tail[-10:]) or ""))


def ensure_dirs():
    os.makedirs(WORDLIB_DIR, exist_ok=True)
    os.makedirs(EXPORT_DIR, exist_ok=True)


def lib_files():
    ensure_dirs()
    result = []
    for fn in os.listdir(WORDLIB_DIR):
        low = fn.lower()
        if low.endswith(".xlsx"):
            result.append(("xlsx", fn))
    return result


def lib_names():
    names = set()
    for _kind, fn in lib_files():
        names.add(os.path.splitext(fn)[0])
    return sorted(names)


def find_lib_path(name):
    for _kind, fn in lib_files():
        if os.path.splitext(fn)[0] == name:
            return os.path.join(WORDLIB_DIR, fn)
    return None


def load_entries(name):
    p = find_lib_path(name)
    if p is None:
        return []
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    entries = []
    for r in rows[1:]:
        if not r:
            continue
        word = "" if r[0] is None else str(r[0]).strip()
        meaning = "" if (len(r) < 2 or r[1] is None) else str(r[1]).strip()
        if word or meaning:
            entries.append({"word": word, "meaning": meaning})
    return entries


def save_excel(name, entries):
    path = os.path.join(WORDLIB_DIR, name + ".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "词句库"
    ws.append(["词句", "中文释义"])
    for e in entries:
        ws.append([e["word"], e["meaning"]])
    wb.save(path)
    return path


def create_excel_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "词句库"
    ws.append(["词句", "中文释义"])
    ws.append(["example", "例子"])
    ws.append(["apple", "苹果"])
    ws.append(["", "（在此下方继续添加，词句留空会跳过）"])
    wb.save(path)


def import_lib(src_path):
    ensure_dirs()
    if not src_path.lower().endswith(".xlsx"):
        raise ValueError("仅支持导入 Excel (.xlsx) 词句库。")
    base = os.path.basename(src_path)
    name = os.path.splitext(base)[0]
    dest = os.path.join(WORDLIB_DIR, base)
    i = 2
    while os.path.exists(dest):
        stem, ext = os.path.splitext(base)
        dest = os.path.join(WORDLIB_DIR, "%s (%d)%s" % (stem, i, ext))
        i += 1
    shutil.copy2(src_path, dest)
    return os.path.splitext(os.path.basename(dest))[0]


def sanitize_filename(s):
    s = str(s).strip()
    for ch in '\\/:*?"<>|':
        s = s.replace(ch, "_")
    s = re.sub(r"\s+", " ", s).strip()
    return (s[:60] or "item")


def make_zip(src_dir, zip_path):
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _dirs, files in os.walk(src_dir):
            for f in files:
                full = os.path.join(root, f)
                rel = os.path.relpath(full, src_dir)
                z.write(full, rel)


def play_wav(path):
    """播放 WAV; 返回需要保持引用的 Sound 对象(避免被回收中断播放)。"""
    try:
        snd = wx.adv.Sound(path)
        snd.Play(wx.adv.SOUND_ASYNC)
        return snd
    except Exception:
        try:
            os.startfile(path)
        except Exception:
            pass
        return None


# ---------------------------------------------------------------------------
# 音频导出配置对话框(后补界面)
# ---------------------------------------------------------------------------

class VisualSettingsDialog(wx.Dialog):
    def __init__(self, parent, config):
        super().__init__(parent, title="视觉设置", size=(460, 320),
                         style=wx.DEFAULT_DIALOG_STYLE)
        self._config = config
        panel = wx.Panel(self)
        root = wx.BoxSizer(wx.VERTICAL)

        visual_box = wx.StaticBox(panel, label="视觉样式")
        visual = wx.StaticBoxSizer(visual_box, wx.VERTICAL)
        grid = wx.FlexGridSizer(3, 2, 10, 10)
        grid.AddGrowableCol(1, 1)

        grid.Add(wx.StaticText(visual_box, label="窗口背景"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 4)
        self.bg_color = wx.ColourPickerCtrl(visual_box, colour=wx.Colour(*hex_to_rgb(config.bg_color)))
        grid.Add(self.bg_color, 0, wx.EXPAND | wx.ALL, 4)

        grid.Add(wx.StaticText(visual_box, label="主强调色"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 4)
        self.accent_color = wx.ColourPickerCtrl(visual_box, colour=wx.Colour(*hex_to_rgb(config.accent_color)))
        grid.Add(self.accent_color, 0, wx.EXPAND | wx.ALL, 4)

        grid.Add(wx.StaticText(visual_box, label="透明度"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 4)
        self.opacity_slider = wx.Slider(visual_box, value=int((config.window_alpha / 255.0) * 100), minValue=60, maxValue=100)
        grid.Add(self.opacity_slider, 0, wx.EXPAND | wx.ALL, 4)

        visual.Add(grid, 0, wx.EXPAND | wx.ALL, 10)
        root.Add(visual, 0, wx.EXPAND | wx.ALL, 8)

        btns = wx.BoxSizer(wx.HORIZONTAL)
        ok_btn = _apply_button_style(wx.Button(panel, wx.ID_OK, "保存", style=wx.BORDER_NONE))
        cancel_btn = _apply_button_style(wx.Button(panel, wx.ID_CANCEL, "取消", style=wx.BORDER_NONE))
        btns.AddStretchSpacer(1)
        btns.Add(ok_btn, 0, wx.ALL, 5)
        btns.Add(cancel_btn, 0, wx.ALL, 5)
        root.Add(btns, 0, wx.EXPAND | wx.ALL, 8)

        panel.SetSizer(root)
        self.SetClientSize(panel.GetBestSize())


class AudioConfigDialog(wx.Dialog):
    def __init__(self, parent, config):
        super().__init__(parent, title="音频导出配置", size=(560, 700),
                         style=wx.DEFAULT_DIALOG_STYLE)
        self._config = config
        self._app = parent
        panel = wx.Panel(self)
        root = wx.BoxSizer(wx.VERTICAL)

        # ---- 语音生成设置 ----
        gen_box = wx.StaticBox(panel, label="语音生成设置")
        gen = wx.StaticBoxSizer(gen_box, wx.VERTICAL)
        grid = wx.FlexGridSizer(4, 2, 8, 8)
        grid.AddGrowableCol(1, 1)

        grid.Add(self._label(gen_box, "语音风格"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        self.voice_cb = wx.ComboBox(gen_box, style=wx.CB_READONLY)
        self.voice_cb.Set([v[1] for v in VOICES])
        grid.Add(self.voice_cb, 0, wx.EXPAND | wx.ALL, 2)

        grid.Add(self._label(gen_box, "语速"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        speed_row = wx.BoxSizer(wx.HORIZONTAL)
        self.speed_slider = wx.Slider(gen_box, value=int(config.speed * 100),
                                      minValue=50, maxValue=200)
        self.speed_text = wx.StaticText(gen_box, label="%.1fx" % config.speed)
        speed_row.Add(self.speed_slider, 1, wx.EXPAND | wx.ALL, 2)
        speed_row.Add(self.speed_text, 0, wx.ALIGN_CENTER_VERTICAL | wx.LEFT | wx.RIGHT, 10)
        grid.Add(speed_row, 0, wx.EXPAND | wx.ALL, 2)

        grid.Add(self._label(gen_box, "模型质量"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        self.quality_cb = wx.ComboBox(gen_box, choices=["fp32", "fp16", "q8", "q4", "q4f16"],
                                      style=wx.CB_READONLY)
        self.quality_cb.SetValue(config.dtype)
        grid.Add(self.quality_cb, 0, wx.EXPAND | wx.ALL, 2)

        grid.Add(self._label(gen_box, "采样率"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        grid.Add(wx.StaticText(gen_box, label="24000 Hz（模型固定）"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)

        gen.Add(grid, 0, wx.EXPAND | wx.ALL, 8)
        root.Add(gen, 0, wx.EXPAND | wx.ALL, 8)

        # ---- HTML 播放器跟读默认设置 ----
        sr_box = wx.StaticBox(panel, label="HTML 播放器跟读默认设置")
        sr = wx.StaticBoxSizer(sr_box, wx.VERTICAL)
        sr_grid = wx.FlexGridSizer(3, 2, 8, 8)
        sr_grid.AddGrowableCol(1, 1)

        sr_grid.Add(self._label(sr_box, "跟读次数"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        self.repeat_spin = wx.SpinCtrl(sr_box, min=1, max=10, initial=config.repeat_count)
        sr_grid.Add(self.repeat_spin, 0, wx.ALL, 2)

        sr_grid.Add(self._label(sr_box, "间隔方式"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        mode_row = wx.BoxSizer(wx.HORIZONTAL)
        self.rb_seconds = wx.RadioButton(sr_box, label="固定秒数", style=wx.RB_GROUP)
        self.rb_multiple = wx.RadioButton(sr_box, label="音频时长倍数")
        self.rb_seconds.SetValue(config.interval_mode == "seconds")
        self.rb_multiple.SetValue(config.interval_mode == "multiple")
        mode_row.Add(self.rb_seconds, 0, wx.ALIGN_CENTER_VERTICAL | wx.RIGHT, 12)
        mode_row.Add(self.rb_multiple, 0, wx.ALIGN_CENTER_VERTICAL, 0)
        sr_grid.Add(mode_row, 0, wx.ALL, 2)

        sr_grid.Add(self._label(sr_box, "间隔值"), 0, wx.ALIGN_CENTER_VERTICAL | wx.ALL, 2)
        self.interval_spin = wx.SpinCtrlDouble(sr_box, min=0.0, max=120.0,
                                               initial=config.interval_value,
                                               inc=0.1)
        self.interval_spin.SetDigits(1)
        sr_grid.Add(self.interval_spin, 0, wx.ALL, 2)

        sr.Add(sr_grid, 0, wx.EXPAND | wx.ALL, 8)
        root.Add(sr, 0, wx.EXPAND | wx.ALL, 8)

        hint = wx.StaticText(panel, label="提示：跟读设置是网页播放器的默认值，网页中仍可实时调整。")
        hint.Wrap(520)
        root.Add(hint, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)

        # ---- 生成测试 ----
        test_box = wx.StaticBox(panel, label="生成测试（使用上方语音风格/语速/质量设置）")
        test = wx.StaticBoxSizer(test_box, wx.VERTICAL)
        self.test_text = wx.TextCtrl(test_box, value=DEFAULT_TEST_TEXT,
                                     style=wx.TE_MULTILINE, size=(-1, 60))
        test.Add(self.test_text, 1, wx.EXPAND | wx.ALL, 5)
        test_row = wx.BoxSizer(wx.HORIZONTAL)
        test_row.AddStretchSpacer(1)
        self.test_btn = _apply_button_style(wx.Button(test_box, label="生成并播放测试音频", style=wx.BORDER_NONE))
        test_row.Add(self.test_btn, 0, wx.ALL, 5)
        test.Add(test_row, 0, wx.EXPAND | wx.ALL, 5)
        root.Add(test, 0, wx.EXPAND | wx.ALL, 8)

        # ---- 按钮 ----
        btns = wx.BoxSizer(wx.HORIZONTAL)
        ok_btn = _apply_button_style(wx.Button(panel, wx.ID_OK, "确定", style=wx.BORDER_NONE))
        cancel_btn = _apply_button_style(wx.Button(panel, wx.ID_CANCEL, "取消", style=wx.BORDER_NONE))
        btns.AddStretchSpacer(1)
        btns.Add(ok_btn, 0, wx.ALL, 8)
        btns.Add(cancel_btn, 0, wx.ALL, 8)
        root.Add(btns, 0, wx.EXPAND | wx.ALL, 8)

        panel.SetSizer(root)

        # 初始选中当前语音
        idx = 0
        for i, (vid, _lbl) in enumerate(VOICES):
            if vid == config.voice:
                idx = i
                break
        self.voice_cb.SetSelection(idx)

        self.speed_slider.Bind(wx.EVT_SLIDER, self._on_speed)
        self.test_btn.Bind(wx.EVT_BUTTON, self._on_test)
        ok_btn.Bind(wx.EVT_BUTTON, self._on_ok)
        cancel_btn.Bind(wx.EVT_BUTTON, lambda e: self.EndModal(wx.ID_CANCEL))

        self.Centre(wx.BOTH)

    def _label(self, parent, text):
        return wx.StaticText(parent, label=text)

    def _on_speed(self, evt):
        self.speed_text.SetLabel("%.1fx" % (self.speed_slider.GetValue() / 100.0))
        evt.Skip()

    def _on_test(self, evt):
        text = self.test_text.GetValue().strip()
        if not text:
            wx.MessageBox("请输入测试文本。", "提示", wx.OK | wx.ICON_INFORMATION)
            return
        voice = VOICES[self.voice_cb.GetSelection()][0]
        speed = self.speed_slider.GetValue() / 100.0
        dtype = self.quality_cb.GetValue()
        out_dir = tempfile.mkdtemp(prefix="zerolisten_test_")
        items = [{"filename": "test.wav", "text": text}]
        ok = self._app._run_tts_export(items, out_dir, "测试音频",
                                       voice=voice, speed=speed, dtype=dtype,
                                       parent=self)
        if ok:
            self._test_sound = play_wav(os.path.join(out_dir, "test.wav"))

    def _on_ok(self, evt):
        self._config.voice = VOICES[self.voice_cb.GetSelection()][0]
        self._config.speed = self.speed_slider.GetValue() / 100.0
        self._config.dtype = self.quality_cb.GetValue()
        self._config.repeat_count = self.repeat_spin.GetValue()
        self._config.interval_mode = "seconds" if self.rb_seconds.GetValue() else "multiple"
        self._config.interval_value = float(self.interval_spin.GetValue())
        self._config._normalize()
        save_config(self._config)
        self.EndModal(wx.ID_OK)


# ---------------------------------------------------------------------------
# 词库编辑页(复用 wordlib_grid, 追加工具按钮)
# ---------------------------------------------------------------------------

class EditorPage(wx.Panel):
    def __init__(self, parent, app):
        super().__init__(parent)
        self.app = app
        root = wx.BoxSizer(wx.VERTICAL)
        self.view = wordlib_grid(self)  # 含词库路径显示 + 表格
        root.Add(self.view, 1, wx.EXPAND | wx.ALL, 5)

        tools = wx.BoxSizer(wx.HORIZONTAL)
        self.btn_excel = _apply_button_style(wx.Button(self, label="保存为Excel词库", style=wx.BORDER_NONE))
        self.btn_template = _apply_button_style(wx.Button(self, label="生成Excel词库模板", style=wx.BORDER_NONE))
        for b in (self.btn_excel, self.btn_template):
            tools.Add(b, 0, wx.ALL, 8)
        root.Add(tools, 0, wx.EXPAND, 5)
        self.SetSizer(root)


# ---------------------------------------------------------------------------
# 主窗口
# ---------------------------------------------------------------------------

class MainFrame(main_panel):
    def __init__(self):
        super().__init__(None)
        self.SetTitle("零听 ZeroListen — 本地词句 TTS 批量生成")
        ensure_dirs()
        self.config = load_config()
        self.current_lib = None
        self.entries = []

        self._build_pages()
        self._bind()
        apply_visual_theme(self, self.config)
        self.refresh_wordlib_list()

    # ---- 页面搭建 ----
    def _build_pages(self):
        self.welcome = welcome_page(self.workbar)
        self.workbar.AddPage(self.welcome, "欢迎")

        self.editor = EditorPage(self.workbar, self)
        self.workbar.AddPage(self.editor, "词库编辑")

        self.exporter = workbar_page(self.workbar)
        self.workbar.AddPage(self.exporter, "导出")


    def _bind(self):
        self.create_wordlib.Bind(wx.EVT_BUTTON, self.on_create_wordlib)
        self.import_wordlib.Bind(wx.EVT_BUTTON, self.on_import_wordlib)
        self.m_listBox1.Bind(wx.EVT_LISTBOX, self.on_select_lib)

        self.welcome.about.Bind(wx.EVT_BUTTON, self.on_about)
        self.welcome.visual_settings.Bind(wx.EVT_BUTTON, self.on_visual_settings)

        self.editor.btn_excel.Bind(wx.EVT_BUTTON, self.on_save_excel)
        self.editor.btn_template.Bind(wx.EVT_BUTTON, self.on_template)


        self.exporter.checkall.Bind(wx.EVT_BUTTON, self.on_checkall)
        self.exporter.output_wav.Bind(wx.EVT_BUTTON, self.on_export_audio)
        self.exporter.output_html.Bind(wx.EVT_BUTTON, self.on_export_html)
        self.exporter.output_audio_config.Bind(wx.EVT_BUTTON, self.on_audio_config)

        self.workbar.Bind(wx.EVT_NOTEBOOK_PAGE_CHANGED, self.on_page_changed)

    # ---- 词句库列表 ----
    def refresh_wordlib_list(self):
        names = lib_names()
        self.m_listBox1.Set(names)
        if self.current_lib in names:
            self.m_listBox1.SetSelection(names.index(self.current_lib))

    def on_select_lib(self, evt):
        name = self.m_listBox1.GetStringSelection()
        if name:
            self.load_lib(name)

    def load_lib(self, name):
        self.entries = load_entries(name)
        self.current_lib = name
        self._load_entries_to_grid()
        self.refresh_checklist()
        self.refresh_wordlib_list()

    def _lib_path_text(self):
        return find_lib_path(self.current_lib) if self.current_lib else ""

    def _load_entries_to_grid(self):
        grid = self.editor.view.grid_wordlib
        n = grid.GetNumberRows()
        if n > 0:
            grid.DeleteRows(0, n)
        grid.AppendRows(len(self.entries) if self.entries else 1)
        for i, e in enumerate(self.entries):
            grid.SetCellValue(i, 0, e["word"])
            grid.SetCellValue(i, 1, e["meaning"])
        self.editor.view.wordlib_path.SetValue(self._lib_path_text())

    def _grid_to_entries(self):
        grid = self.editor.view.grid_wordlib
        out = []
        for i in range(grid.GetNumberRows()):
            word = grid.GetCellValue(i, 0).strip()
            meaning = grid.GetCellValue(i, 1).strip()
            if word or meaning:
                out.append({"word": word, "meaning": meaning})
        return out

    def _sync_entries_from_grid(self):
        """把表格里的最新编辑同步到 entries 并刷新导出勾选列表。"""
        if self.current_lib is None:
            return
        entries = self._grid_to_entries()
        if entries != self.entries:
            self.entries = entries
            self.refresh_checklist()

    def on_page_changed(self, evt):
        if evt.GetSelection() == 2:  # 导出页
            self._sync_entries_from_grid()
        evt.Skip()

    # ---- 新建 / 导入 ----
    def on_create_wordlib(self, evt):
        name = wx.GetTextFromUser("请输入新词句库名称：", "创建新词句库").strip()
        if not name:
            return
        if name in lib_names():
            wx.MessageBox("该词句库已存在。", "提示", wx.OK | wx.ICON_INFORMATION)
            return
        save_excel(name, [])
        self.current_lib = name
        self.entries = []
        self.refresh_wordlib_list()
        self.load_lib(name)

    def on_import_wordlib(self, evt):
        dlg = wx.FileDialog(
            self, "导入词句库", defaultDir=WORDLIB_DIR,
            wildcard="Excel 词库 (*.xlsx)|*.xlsx",
            style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST)
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        src = dlg.GetPath()
        dlg.Destroy()
        try:
            name = import_lib(src)
        except ValueError as ex:
            wx.MessageBox(str(ex), "提示", wx.OK | wx.ICON_INFORMATION)
            return
        self.refresh_wordlib_list()
        self.load_lib(name)

    # ---- 保存 / 模板 ----
    def _require_lib(self):
        if not self.current_lib:
            wx.MessageBox("请先选择或创建一个词句库。", "提示", wx.OK | wx.ICON_INFORMATION)
            return False
        return True

    def on_save_excel(self, evt):
        if not self._require_lib():
            return
        entries = self._grid_to_entries()
        path = save_excel(self.current_lib, entries)
        self.entries = entries
        self.refresh_wordlib_list()
        self.refresh_checklist()
        self.editor.view.wordlib_path.SetValue(path)
        wx.MessageBox("已保存：\n" + path, "完成", wx.OK | wx.ICON_INFORMATION)

    def on_template(self, evt):
        dlg = wx.FileDialog(self, "保存 Excel 词库模板", defaultDir=WORDLIB_DIR,
                            defaultFile="词句库模板.xlsx", wildcard="Excel (*.xlsx)|*.xlsx",
                            style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        p = dlg.GetPath()
        dlg.Destroy()
        create_excel_template(p)
        try:
            os.startfile(os.path.dirname(p))
        except Exception:
            pass
        wx.MessageBox("已生成模板：\n" + p, "完成", wx.OK | wx.ICON_INFORMATION)

    # ---- 导出页 ----
    def refresh_checklist(self):
        cb = self.exporter.check_words
        cb.Clear()
        for e in self.entries:
            label = e["word"]
            if e["meaning"]:
                label += "　" + e["meaning"]
            cb.Append(label)
        for i in range(cb.GetCount()):
            cb.Check(i, True)

    def on_checkall(self, evt):
        cb = self.exporter.check_words
        if cb.GetCount() == 0:
            return
        any_checked = any(cb.IsChecked(i) for i in range(cb.GetCount()))
        for i in range(cb.GetCount()):
            cb.Check(i, not any_checked)

    def selected_entries(self):
        cb = self.exporter.check_words
        sel = []
        for i in cb.GetCheckedItems():
            if 0 <= i < len(self.entries):
                sel.append(self.entries[i])
        return sel

    def on_audio_config(self, evt):
        dlg = AudioConfigDialog(self, self.config)
        dlg.ShowModal()
        dlg.Destroy()

    def on_visual_settings(self, evt):
        dlg = VisualSettingsDialog(self, self.config)
        if dlg.ShowModal() == wx.ID_OK:
            self.config.bg_color = dlg.bg_color.GetColour().GetAsString(wx.C2S_HTML_SYNTAX)
            self.config.accent_color = dlg.accent_color.GetColour().GetAsString(wx.C2S_HTML_SYNTAX)
            self.config.window_alpha = int(round((dlg.opacity_slider.GetValue() / 100.0) * 255))
            self.config._normalize()
            save_config(self.config)
            apply_visual_theme(self, self.config)
        dlg.Destroy()

    def on_export_audio(self, evt):
        self._sync_entries_from_grid()
        sel = self.selected_entries()
        if not sel:
            wx.MessageBox("请先选择要导出的词句。", "提示", wx.OK | wx.ICON_INFORMATION)
            return
        dlg = wx.DirDialog(self, "选择音频导出目录", defaultPath=EXPORT_DIR)
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        out_dir = dlg.GetPath()
        dlg.Destroy()
        os.makedirs(out_dir, exist_ok=True)

        items = []
        for idx, e in enumerate(sel):
            fname = "%04d_%s.wav" % (idx + 1, sanitize_filename(e["word"]))
            items.append({"filename": fname, "text": e["word"]})
        if self._run_tts_export(items, out_dir, "音频"):
            wx.MessageBox("已导出 %d 条音频到：\n%s" % (len(items), out_dir),
                          "完成", wx.OK | wx.ICON_INFORMATION)

    def on_export_html(self, evt):
        self._sync_entries_from_grid()
        sel = self.selected_entries()
        if not sel:
            wx.MessageBox("请先选择要导出的词句。", "提示", wx.OK | wx.ICON_INFORMATION)
            return
        base = self.current_lib or "wordlib"
        dlg = wx.FileDialog(self, "保存 HTML 播放器压缩包", defaultDir=EXPORT_DIR,
                            defaultFile=base + "_播放器.zip", wildcard="ZIP 压缩包 (*.zip)|*.zip",
                            style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT)
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        zip_path = dlg.GetPath()
        dlg.Destroy()

        tmp = tempfile.mkdtemp(prefix="zerolisten_")
        audio_dir = os.path.join(tmp, "audio")
        os.makedirs(audio_dir, exist_ok=True)

        items = []
        for idx, e in enumerate(sel):
            items.append({"filename": "%04d.wav" % (idx + 1), "text": e["word"]})

        try:
            if not self._run_tts_export(items, audio_dir, "HTML 播放器音频"):
                return
            playlist = [
                {"word": e["word"], "meaning": e["meaning"],
                 "audio": "audio/%04d.wav" % (i + 1)}
                for i, e in enumerate(sel)
            ]
            html = render_player_html(playlist, title=(self.current_lib or "ZeroListen") + " 播放器",
                                      config=self.config.to_dict())
            with open(os.path.join(tmp, "index.html"), "w", encoding="utf-8") as f:
                f.write(html)
            make_zip(tmp, zip_path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

        wx.MessageBox("已导出 HTML 播放器：\n" + zip_path +
                      "\n解压后打开 index.html 即可播放。", "完成", wx.OK | wx.ICON_INFORMATION)

    # ---- TTS 执行 ----
    def _run_tts_export(self, items, out_dir, label, voice=None, speed=None, dtype=None, parent=None):
        """生成音频(带进度对话框), 成功返回 True。

        voice/speed/dtype 缺省时使用已保存的配置; 可为单条测试临时覆盖。
        """
        voice = voice or self.config.voice
        if speed is None:
            speed = self.config.speed
        dtype = dtype or self.config.dtype

        full_items = []
        for it in items:
            full_items.append({
                "filename": it["filename"],
                "text": it["text"],
                "voice": it.get("voice") or voice,
                "speed": it.get("speed") if it.get("speed") is not None else speed,
            })
        manifest = {
            "model_id": DEFAULT_MODEL_ID,
            "dtype": dtype,
            "output_dir": out_dir,
            "items": full_items,
        }
        tmp = tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8")
        json.dump(manifest, tmp, ensure_ascii=False)
        tmp.close()

        dlg = progress(parent if parent is not None else self)
        dlg.progress_text.SetLabel("正在生成" + label)
        dlg.info.SetLabel("正在准备语音模型（首次运行需下载，之后使用缓存）…")
        start = time.time()
        result = {"ok": False, "error": ""}

        def prog(i, n, label=None):
            pct = int(round(i * 100.0 / n)) if n else 0
            elapsed = time.time() - start
            eta = (elapsed / i * (n - i)) if i > 0 else 0.0
            if label:
                msg = label
            else:
                msg = "正在进行第%d项，预计还需 %.0f 秒\n已完成 %d%% (%d/%d)" % (i, eta, pct, i, n)

            def update():
                dlg.m_gauge1.SetValue(min(100, max(0, pct)))
                dlg.info.SetLabel(msg)
            wx.CallAfter(update)

        def worker():
            try:
                self._run_node(tmp.name, prog)
                result["ok"] = True
                wx.CallAfter(dlg.EndModal, wx.ID_OK)
            except Exception as ex:
                result["error"] = str(ex)
                wx.CallAfter(dlg.EndModal, wx.ID_CANCEL)

        threading.Thread(target=worker, daemon=True).start()
        dlg.ShowModal()
        dlg.Destroy()

        try:
            os.remove(tmp.name)
        except OSError:
            pass

        if not result["ok"]:
            show_error_dialog("生成失败", result["error"])
            return False
        return True

    def _run_node(self, manifest_path, prog_cb):
        node = _resolve_node_executable(prog_cb)
        if not node:
            raise RuntimeError("未找到 Node.js，且自动下载失败。请确认网络正常，或手动放置 node.exe 到项目根目录的 node 文件夹。")
        _ensure_node_dependencies(node, prog_cb)
        cmd = [node, TTS_CLI, manifest_path]
        proc = subprocess.Popen(cmd, cwd=APP_DIR, stdout=subprocess.PIPE,
                                stderr=subprocess.STDOUT, text=True,
                                encoding="utf-8", errors="replace")
        tail = []
        for line in proc.stdout:
            line = line.rstrip("\n")
            tail.append(line)
            if len(tail) > 40:
                tail.pop(0)
            m = re.match(r"PROGRESS (\d+)/(\d+)", line)
            if m:
                prog_cb(int(m.group(1)), int(m.group(2)))
        rc = proc.wait()
        if rc != 0:
            raise RuntimeError("\n".join(tail[-10:]) or ("node 退出码 %d" % rc))

    # ---- 欢迎页 ----
    def on_placeholder(self, evt):
        wx.MessageBox("该功能尚未开放，敬请期待。", "提示", wx.OK | wx.ICON_INFORMATION)

    def on_about(self, evt):
        wx.MessageBox("零听 ZeroListen\n本地离线中/英文词句 TTS 批量生成工具\n\n"
                      "依赖本地 ONNX 模型与 kokoro-js，\n不调用云端语音接口。\nhttps://github.com/Loser123zbx",
                      "关于", wx.OK | wx.ICON_INFORMATION)


def main():
    install_global_exception_handler()
    app = wx.App(False)
    frame = MainFrame()
    frame.Show()
    app.MainLoop()


if __name__ == "__main__":
    main()
